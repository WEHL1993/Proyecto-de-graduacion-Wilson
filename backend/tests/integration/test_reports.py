"""Reportes gerenciales contra PostgreSQL real (`/reports`, permiso `reportes:leer`).

Los datos usan fechas de 2031 y filtran por ruta para no mezclarse con el histórico real de la BD.
"""

import io
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.domain.models.catalog import Categoria, Producto, Ruta
from app.domain.models.ml import ModeloML, PronosticoDemanda
from app.domain.models.operations import CargaRuta, DetalleCarga
from app.domain.models.sales import Comision, EtlLote, VentaHistorica

from .conftest import crear_usuario, encabezados

BASE = "/api/v1/reports"
D10, D11, D12, D13 = (date(2031, 1, d) for d in (10, 11, 12, 13))
RANGO = {"desde": "2031-01-10", "hasta": "2031-01-13"}


@pytest.fixture
def entorno(db_tx, cliente):
    db = db_tx
    sufijo = uuid.uuid4().hex[:6]
    gerente = crear_usuario(db, "gerente")
    vendedor = crear_usuario(db, "Vendedor Uno")
    categoria = Categoria(nombre=f"cat-{sufijo}")
    ruta1 = Ruta(codigo=f"R1-{sufijo}", nombre="Ruta Uno", vendedor_id=vendedor.id)
    ruta2 = Ruta(codigo=f"R2-{sufijo}", nombre="Ruta Dos", vendedor_id=vendedor.id)
    modelo = ModeloML(
        nombre=f"demanda-{sufijo}",
        algoritmo="xgboost",
        version="v1.0.1",
        ruta_artefacto="models/xgboost/v1.0.1",
        hash_artefacto="0" * 64,
        hiperparametros={},
        esquema_features=[],
        ventana_desde=date(2025, 1, 1),
        ventana_hasta=date(2025, 3, 31),
        motivo_entrenamiento="manual",
    )
    lote = EtlLote(
        usuario_id=gerente.id, archivo_nombre="reportes.xlsx", checksum_sha256=uuid.uuid4().hex * 2
    )
    db.add_all([categoria, ruta1, ruta2, modelo, lote])
    db.flush()
    prod_a = Producto(
        sku=f"A-{sufijo}", nombre="A", categoria_id=categoria.id, costo_unitario=Decimal("2")
    )
    prod_b = Producto(
        sku=f"B-{sufijo}", nombre="B", categoria_id=categoria.id, costo_unitario=Decimal("4")
    )
    db.add_all([prod_a, prod_b])
    db.flush()

    def venta(fecha, producto, ruta, cantidad, precio):
        return VentaHistorica(
            lote_id=lote.id,
            fecha_venta=fecha,
            producto_id=producto.id,
            ruta_id=ruta.id,
            vendedor_id=vendedor.id,
            cantidad=Decimal(cantidad),
            precio_unitario=Decimal(precio),
            monto_total=Decimal(cantidad) * Decimal(precio),
        )

    # Ruta 1: 35 uds, monto 190, costo 80. Ruta 2: 8 uds, monto 40, costo 16.
    ventas = [
        venta(D10, prod_a, ruta1, "10", "5"),
        venta(D10, prod_b, ruta1, "5", "8"),
        venta(D11, prod_a, ruta1, "20", "5"),
        venta(D10, prod_a, ruta2, "8", "5"),
    ]
    db.add_all(ventas)
    db.flush()
    db.add_all(
        Comision(
            vendedor_id=vendedor.id,
            venta_id=v.id,
            periodo="2031-01",
            porcentaje=Decimal("3"),
            monto=v.monto_total * Decimal("0.03"),
        )
        for v in ventas
    )

    def carga(fecha, estado):
        extra = (
            {"aprobado_por": gerente.id, "aprobada_en": datetime.now(UTC)}
            if estado == "aprobada"
            else {}
        )
        return CargaRuta(
            ruta_id=ruta1.id,
            fecha_operacion=fecha,
            estado=estado,
            modelo_id=modelo.id,
            generado_por=gerente.id,
            **extra,
        )

    def detalle(c, producto, predicha, disponible, sugerida, ajustado):
        return DetalleCarga(
            carga_id=c.id,
            producto_id=producto.id,
            cantidad_predicha=Decimal(predicha),
            stock_disponible_al_generar=Decimal(disponible),
            cantidad_sugerida=Decimal(sugerida),
            ajustado_por_stock=ajustado,
        )

    aprobada, rechazada = carga(D12, "aprobada"), carga(D13, "rechazada")
    db.add_all([aprobada, rechazada])
    db.flush()
    db.add_all(
        [
            detalle(aprobada, prod_a, "30", "20", "20", True),  # 10 uds sin cubrir
            detalle(aprobada, prod_b, "10", "50", "10", False),
            detalle(rechazada, prod_a, "30", "20", "20", True),  # excluida: rechazada
        ]
    )

    def proyeccion(fecha, producto, demanda, horizonte, generado):
        return PronosticoDemanda(
            modelo_id=modelo.id,
            producto_id=producto.id,
            ruta_id=ruta1.id,
            fecha_objetivo=fecha,
            horizonte_dias=horizonte,
            demanda_predicha=Decimal(demanda),
            generado_en=generado,
        )

    db.add_all(
        [
            proyeccion(D10, prod_a, "12", 1, datetime(2031, 1, 1, tzinfo=UTC)),
            proyeccion(D10, prod_a, "14", 2, datetime(2031, 1, 5, tzinfo=UTC)),  # la vigente
            proyeccion(D11, prod_b, "6", 1, datetime(2031, 1, 1, tzinfo=UTC)),
        ]
    )
    db.commit()
    return SimpleNamespace(
        db=db,
        client=cliente,
        ruta1=ruta1,
        ruta2=ruta2,
        vendedor=vendedor,
        h=encabezados(gerente, "reportes:leer", roles=("Gerente",)),
        gerente=gerente,
    )


def test_exige_permiso_y_autenticacion(entorno):
    e = entorno
    for ruta in ("inventory-turnover", "commissions", "sales-vs-forecast", "export"):
        assert e.client.get(f"{BASE}/{ruta}").status_code == 401
        sin = encabezados(e.gerente, "alertas:leer")
        assert e.client.get(f"{BASE}/{ruta}", headers=sin).status_code == 403


def test_rotacion_e_indice_de_quiebres_por_ruta(entorno):
    e = entorno
    r = e.client.get(
        f"{BASE}/inventory-turnover", params={**RANGO, "ruta_id": str(e.ruta1.id)}, headers=e.h
    )
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert len(cuerpo["rutas"]) == 1
    fila = cuerpo["rutas"][0]
    assert fila["ruta_id"] == str(e.ruta1.id)
    assert Decimal(fila["unidades_vendidas"]) == 35 and Decimal(fila["monto_vendido"]) == 190
    assert Decimal(fila["costo_ventas"]) == 80
    # La carga rechazada no cuenta: 2 líneas, 1 ajustada por stock → 50 %.
    assert (fila["cargas"], fila["lineas_carga"], fila["lineas_ajustadas"]) == (1, 2, 1)
    assert Decimal(fila["indice_quiebre"]) == 50 and Decimal(fila["unidades_no_cubiertas"]) == 10
    assert Decimal(cuerpo["costo_ventas_total"]) == 80
    if Decimal(cuerpo["valor_inventario"]) > 0:
        esperado = Decimal(80) / Decimal(cuerpo["valor_inventario"])
        assert abs(Decimal(cuerpo["rotacion_global"]) - esperado) < Decimal("0.001")
    else:
        assert cuerpo["rotacion_global"] is None and cuerpo["dias_inventario"] is None


def test_rotacion_de_dos_rutas_ordena_por_costo(entorno):
    e = entorno
    r = e.client.get(f"{BASE}/inventory-turnover", params=RANGO, headers=e.h).json()
    por_id = {f["ruta_id"]: f for f in r["rutas"]}
    assert Decimal(por_id[str(e.ruta2.id)]["costo_ventas"]) == 16
    assert por_id[str(e.ruta2.id)]["indice_quiebre"] is None  # sin cargas
    ids = [f["ruta_id"] for f in r["rutas"]]
    assert ids.index(str(e.ruta1.id)) < ids.index(str(e.ruta2.id))


def test_liquidacion_de_comisiones(entorno):
    e = entorno
    params = {"periodo_desde": "2031-01", "periodo_hasta": "2031-01"}
    r = e.client.get(
        f"{BASE}/commissions", params={**params, "vendedor_id": str(e.vendedor.id)}, headers=e.h
    )
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert Decimal(cuerpo["total_vendido"]) == 230 and Decimal(
        cuerpo["total_comisiones"]
    ) == Decimal("6.90")
    (liq,) = cuerpo["liquidaciones"]
    assert liq["vendedor"] == "Vendedor Uno" and liq["periodo"] == "2031-01"
    assert liq["ventas_registradas"] == 4 and Decimal(liq["porcentaje_efectivo"]) == 3

    solo_ruta1 = e.client.get(
        f"{BASE}/commissions",
        params={**params, "vendedor_id": str(e.vendedor.id), "ruta_id": str(e.ruta1.id)},
        headers=e.h,
    ).json()
    assert Decimal(solo_ruta1["total_vendido"]) == 190
    assert Decimal(solo_ruta1["total_comisiones"]) == Decimal("5.70")


def test_validaciones_de_periodo_y_rango(entorno):
    e = entorno
    mal = e.client.get(f"{BASE}/commissions", params={"periodo_desde": "2031-13"}, headers=e.h)
    assert mal.status_code == 422
    invertido = e.client.get(
        f"{BASE}/commissions",
        params={"periodo_desde": "2031-03", "periodo_hasta": "2031-01"},
        headers=e.h,
    )
    assert invertido.status_code == 400 and invertido.json()["codigo"] == "RANGO_INVALIDO"
    fechas = e.client.get(
        f"{BASE}/sales-vs-forecast",
        params={"desde": "2031-02-01", "hasta": "2031-01-01"},
        headers=e.h,
    )
    assert fechas.status_code == 400 and fechas.json()["codigo"] == "RANGO_INVALIDO"
    inexistente = e.client.get(
        f"{BASE}/sales-vs-forecast", params={**RANGO, "ruta_id": str(uuid.uuid4())}, headers=e.h
    )
    assert inexistente.status_code == 404


def test_venta_real_vs_proyectada_usa_la_proyeccion_mas_reciente(entorno):
    e = entorno
    r = e.client.get(
        f"{BASE}/sales-vs-forecast", params={**RANGO, "ruta_id": str(e.ruta1.id)}, headers=e.h
    )
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    (fila,) = cuerpo["rutas"]
    # Proyectada = 14 (horizonte 2, más reciente que el de 12) + 6 = 20; real = 35.
    assert Decimal(fila["real"]) == 35 and Decimal(fila["proyectada"]) == 20
    assert Decimal(fila["desviacion_pct"]) == 75
    assert Decimal(cuerpo["total_real"]) == 35 and Decimal(cuerpo["total_proyectado"]) == 20
    dias = {d["fecha"]: (Decimal(d["real"]), Decimal(d["proyectada"])) for d in cuerpo["serie"]}
    assert dias == {"2031-01-10": (15, 14), "2031-01-11": (20, 6)}

    sin_proyeccion = e.client.get(f"{BASE}/sales-vs-forecast", params=RANGO, headers=e.h).json()
    ruta2 = next(f for f in sin_proyeccion["rutas"] if f["ruta_id"] == str(e.ruta2.id))
    assert Decimal(ruta2["real"]) == 8 and Decimal(ruta2["proyectada"]) == 0
    assert ruta2["desviacion_pct"] is None


def test_exportacion_excel_consolidado_y_csv(entorno):
    e = entorno
    params = {**RANGO, "periodo_desde": "2031-01", "periodo_hasta": "2031-01"}
    xlsx = e.client.get(f"{BASE}/export", params=params, headers=e.h)
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="reporte_consolidado_' in xlsx.headers["content-disposition"]
    libro = load_workbook(io.BytesIO(xlsx.content))
    assert libro.sheetnames == ["Rotación y quiebres", "Comisiones", "Real vs proyectado"]
    comisiones = [f for f in libro["Comisiones"].iter_rows(values_only=True)]
    assert comisiones[0][:2] == ("Periodo", "Vendedor")
    assert any(f[1] == "Vendedor Uno" for f in comisiones[1:])

    csv = e.client.get(
        f"{BASE}/export", params={**params, "tipo": "comisiones", "formato": "csv"}, headers=e.h
    )
    assert csv.status_code == 200 and csv.headers["content-type"].startswith("text/csv")
    texto = csv.content.decode("utf-8-sig")
    assert texto.splitlines()[0] == "Periodo,Vendedor,Ventas,Monto vendido,Comisión,% efectivo"
    assert "Vendedor Uno" in texto

    consolidado_csv = e.client.get(
        f"{BASE}/export", params={**params, "formato": "csv"}, headers=e.h
    ).content.decode("utf-8-sig")
    assert "Rotación y quiebres" in consolidado_csv and "Real vs proyectado" in consolidado_csv
