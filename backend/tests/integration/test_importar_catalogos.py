"""Importador de catálogos (M02) contra PostgreSQL real, con un libro sintético en `tmp_path`.

No se usa ningún Excel real ni el de comisiones: el libro de prueba se construye aquí.
"""

import csv
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select

from app.db.seeds import importar_catalogos as imp
from app.domain.models.catalog import (
    Categoria,
    Inventario,
    Producto,
    ProductoAliasExcel,
    Ruta,
)

PRODUCTOS = [
    # (código, descripción, precio)
    (900001, "BIG COLA PET NO RETORNABLE 3030 ML 6", 56),
    (900002, "BIG COLA PET NO RETORNABLE 3300 ml 6 pack", 60),
    (900003, "VOLT GO PET NO RETORNABLE 625 ml 12 pack", 46.8),
    (900004, "CIFRUT FRESA DISPLAY 70 GRS 12 MC", 40),
    (900005, "DEPORADE UVA PET NO RETORNABLE 1.5ML 6", 40),
]


@pytest.fixture
def libro(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    precios = wb.create_sheet("CODIGOS Y PRECIOS")
    precios.append(["Cantidad", "CODIGO", "DESCRIPCION", "PRECIO", "efectivo"])
    for codigo, descripcion, precio in PRODUCTOS:
        precios.append([None, codigo, descripcion, precio, 0])
    precios.append([None, 900001, "BIG COLA PET NO RETORNABLE 3030 ML 6", 56, 0])  # código repetido
    precios.append([None, None, None, None, 0])  # fila vacía con efectivo
    precios.append([None, 900099, "SIN PRECIO PET 500 ML 6", None, 0])  # precio ausente

    for nombre, filas in (
        ("Cornelio", [(3030, "Cola"), (3300, "Cola"), (625, "Go"), (3030, "Misterioso")]),
        ("Antonio", [(3030, "Cola"), (1.5, "Fruty")]),
        ("CONSUMO", [(3030, "Cola")]),
        ("Merma", [(1, "Basura")]),
        ("Hoja Rara", [(1, "X")]),
    ):
        hoja = wb.create_sheet(nombre)
        hoja.append(["Articulo", "Descripcion", "Saldo 31"])
        for medida, sabor in filas:
            hoja.append([medida, sabor, 1])

    destino = tmp_path / "VENTAS DIARIAS PRUEBA.xlsx"
    wb.save(destino)
    return destino


def _conteos(db) -> tuple[int, ...]:
    return tuple(
        db.scalar(select(func.count()).select_from(m))
        for m in (Producto, Inventario, Categoria, Ruta, ProductoAliasExcel)
    )


def _alias_csv(tmp_path: Path, filas: list[tuple[str, str]]) -> Path:
    destino = tmp_path / "alias.csv"
    with destino.open("w", newline="", encoding="utf-8") as fh:
        escritor = csv.writer(fh)
        escritor.writerow(["alias", "sku"])
        escritor.writerows(filas)
    return destino


def _fila(db, sku: str) -> Producto:
    return db.scalar(select(Producto).where(Producto.sku == sku))


# ------------------------------------------------------------------ dry-run
def test_dry_run_no_modifica_la_base(db_tx, libro, tmp_path):
    antes = _conteos(db_tx)
    r = imp.ejecutar(db_tx, libro, tmp_path / "no-existe.csv", aplicar=False)
    assert _conteos(db_tx) == antes
    assert r.aplicado is False
    assert len(r.productos_creados) == 5 and len(r.rutas_creadas) == 2
    assert sorted(r.categorias_creadas) == ["Big", "Cifrut", "Deporade", "Volt"]


# ------------------------------------------------------------------ aplicar e idempotencia
def test_aplicar_crea_productos_con_presentacion_en_paquete(db_tx, libro, tmp_path):
    imp.ejecutar(db_tx, libro, None, aplicar=True)
    p = _fila(db_tx, "900002")
    assert (p.medida_ml, p.unidades_por_paquete, p.sabor) == (3300, 6, "COLA")
    assert p.unidad_medida == "paquete" and p.precio_venta == Decimal("60.00")
    assert p.proveedor_id is None and p.stock_minimo == 0 and p.activo is True
    assert _fila(db_tx, "900003").precio_venta == Decimal("46.80")
    assert _fila(db_tx, "900004").medida_ml is None  # 70 GRS: sin medida en ml
    assert _fila(db_tx, "900005").medida_ml == 1500  # 1.5ML
    inv = db_tx.scalar(select(Inventario).where(Inventario.producto_id == p.id))
    assert inv is not None and inv.stock_actual == 0


def test_aplicar_crea_solo_las_rutas_de_vendedor_sin_vendedor_asignado(db_tx, libro):
    r = imp.ejecutar(db_tx, libro, None, aplicar=True)
    codigos = set(db_tx.scalars(select(Ruta.codigo)))
    assert {"CORNELIO", "ANTONIO"} <= codigos
    assert not {"CONSUMO", "MERMA", "HOJA RARA"} & codigos
    assert all(x.vendedor_id is None for x in db_tx.scalars(select(Ruta)))
    assert set(r.hojas_ignoradas) == {"CONSUMO", "Merma", "CODIGOS Y PRECIOS"}
    assert r.hojas_no_clasificadas == ["Hoja Rara"]


def test_dos_corridas_dan_el_mismo_resultado(db_tx, libro, tmp_path):
    imp.ejecutar(db_tx, libro, None, aplicar=True)
    tras_primera = _conteos(db_tx)
    r2 = imp.ejecutar(db_tx, libro, None, aplicar=True)
    assert _conteos(db_tx) == tras_primera
    assert r2.productos_creados == [] and r2.productos_actualizados == []
    assert len(r2.productos_sin_cambios) == 5
    assert r2.rutas_creadas == [] and r2.categorias_creadas == []


def test_un_cambio_de_precio_se_actualiza_y_conserva_categoria_y_stock(db_tx, libro):
    imp.ejecutar(db_tx, libro, None, aplicar=True)
    p = _fila(db_tx, "900001")
    otra = Categoria(nombre="Manual")
    db_tx.add(otra)
    db_tx.flush()
    p.categoria_id = otra.id
    p.precio_venta = Decimal("1.00")
    p.costo_unitario = Decimal("9.00")
    db_tx.flush()

    r = imp.ejecutar(db_tx, libro, None, aplicar=True)
    assert r.productos_actualizados == ["900001"]
    db_tx.refresh(p)
    assert p.precio_venta == Decimal("56.00")
    assert p.categoria_id == otra.id and p.costo_unitario == Decimal("9.00")


def test_el_reporte_lista_filas_no_interpretables_duplicados_y_avisos(db_tx, libro):
    r = imp.ejecutar(db_tx, libro, None, aplicar=False)
    motivos = " | ".join(d for _, d in r.no_interpretables)
    assert "código ausente" in motivos and "precio ausente" in motivos
    assert any("900001" in d and "repetido" in d for _, d in r.duplicados)
    avisos = {(sku, a.split(":")[0]) for sku, a in r.avisos_producto}
    assert ("900004", "sin_medida_ml") in avisos and ("900005", "medida_en_litros") in avisos


# ------------------------------------------------------------------ alias y ventas sin empate
def test_sin_csv_de_alias_se_omite_sin_error_y_se_listan_las_ventas_sin_empate(
    db_tx, libro, tmp_path
):
    r = imp.ejecutar(db_tx, libro, tmp_path / "no-existe.csv", aplicar=True)
    assert r.alias_creados == []
    sin_empate = {clave for clave, _ in r.sin_empate}
    assert sin_empate == {"3030-COLA", "3300-COLA", "625-GO", "3030-MISTERIOSO", "1.5-FRUTY"}
    assert any("omite" in n for n in r.notas)


def test_alias_validos_empatan_y_el_resto_se_reporta(db_tx, libro, tmp_path):
    csv_alias = _alias_csv(
        tmp_path,
        [
            ("3030-COLA", "900001"),
            ("3300-COLA", "900002"),
            ("625-GO", "900003"),
            ("1.5-FRUTY", "900005"),
            ("3030-MISTERIOSO", "999999"),  # SKU inexistente
        ],
    )
    r = imp.ejecutar(db_tx, libro, csv_alias, aplicar=True)
    assert sorted(r.alias_creados) == ["1.5-FRUTY", "3030-COLA", "3300-COLA", "625-GO"]
    assert [a for a, _ in r.alias_invalidos] == ["3030-MISTERIOSO"]
    assert {c for c, _ in r.sin_empate} == {"3030-MISTERIOSO"}
    alias = db_tx.scalar(select(ProductoAliasExcel).where(ProductoAliasExcel.alias == "3030-COLA"))
    assert alias.producto_id == _fila(db_tx, "900001").id


def test_alias_idempotente_y_con_conflictos_reportados(db_tx, libro, tmp_path):
    csv_alias = _alias_csv(tmp_path, [("3030-COLA", "900001")])
    imp.ejecutar(db_tx, libro, csv_alias, aplicar=True)
    r = imp.ejecutar(db_tx, libro, csv_alias, aplicar=True)
    assert r.alias_creados == [] and r.alias_sin_cambios == ["3030-COLA"]

    # Mismo alias apuntando a otro SKU que el ya guardado: no se pisa, se reporta.
    conflicto = _alias_csv(tmp_path, [("3030-COLA", "900002")])
    r = imp.ejecutar(db_tx, libro, conflicto, aplicar=True)
    assert [a for a, _ in r.alias_ambiguos] == ["3030-COLA"]
    alias = db_tx.scalar(select(ProductoAliasExcel).where(ProductoAliasExcel.alias == "3030-COLA"))
    assert alias.producto_id == _fila(db_tx, "900001").id


def test_alias_que_apunta_a_mas_de_un_producto_se_reporta_y_no_se_crea(db_tx, libro, tmp_path):
    csv_alias = _alias_csv(tmp_path, [("625-GO", "900003"), ("625-GO", "900002")])
    r = imp.ejecutar(db_tx, libro, csv_alias, aplicar=True)
    assert [a for a, _ in r.alias_ambiguos] == ["625-GO"]
    assert "900002" in r.alias_ambiguos[0][1] and "900003" in r.alias_ambiguos[0][1]
    assert db_tx.scalar(select(func.count()).select_from(ProductoAliasExcel)) == 0


def test_alias_en_dry_run_valida_contra_el_catalogo_planificado_sin_escribir(
    db_tx, libro, tmp_path
):
    csv_alias = _alias_csv(tmp_path, [("3030-COLA", "900001")])
    r = imp.ejecutar(db_tx, libro, csv_alias, aplicar=False)
    assert r.alias_creados == ["3030-COLA"] and r.alias_invalidos == []
    assert db_tx.scalar(select(func.count()).select_from(ProductoAliasExcel)) == 0


# ------------------------------------------------------------------ seguridad y reporte
def test_el_excel_de_comisiones_se_rechaza(tmp_path):
    with pytest.raises(ValueError, match="sensibles"):
        imp._abrir(tmp_path / "COMISIONES VENTAS MARZO 2025.xlsx")


def test_el_csv_de_reporte_no_contiene_datos_fuera_del_catalogo(db_tx, libro, tmp_path):
    r = imp.ejecutar(db_tx, libro, None, aplicar=False)
    destino = imp.escribir_csv(r, tmp_path / "reportes")
    contenido = destino.read_text(encoding="utf-8-sig")
    assert destino.parent.name == "reportes" and contenido.startswith("modo,seccion")
    assert "dry-run,productos_creados,900001" in contenido
    for prohibido in ("nit", "password", "contraseña", "usuario:"):
        assert prohibido not in contenido.lower()


def test_los_casos_reales_del_encargo(db_tx, libro):
    imp.ejecutar(db_tx, libro, None, aplicar=True)
    assert (_fila(db_tx, "900001").medida_ml, _fila(db_tx, "900001").unidades_por_paquete) == (
        3030,
        6,
    )
    assert (_fila(db_tx, "900002").medida_ml, _fila(db_tx, "900002").unidades_por_paquete) == (
        3300,
        6,
    )
    assert (_fila(db_tx, "900003").medida_ml, _fila(db_tx, "900003").unidades_por_paquete) == (
        625,
        12,
    )
    assert _fila(db_tx, "900004").medida_ml is None
