"""Importador idempotente de catálogos desde los Excel de la empresa (M02, ADR-18).

Uso (desde `backend/`, con el venv activo):
    python -m app.db.seeds.importar_catalogos            # dry-run (no escribe nada)
    python -m app.db.seeds.importar_catalogos --aplicar  # escribe (upsert por SKU)

Qué lee (solo lectura, `data/raw/` es inmutable):
- Libro «VENTAS DIARIAS*.xlsx», hoja `CODIGOS Y PRECIOS`: código (= SKU), descripción y precio.
- Del mismo libro, los nombres de las hojas de vendedor (rutas) y sus columnas `Articulo` y
  `Descripcion`, solo para detectar productos que no empatan con ningún SKU ni alias.
- Opcional `data/plantillas/alias_excel.csv` (columnas `alias`, `sku`): se omite si no existe.

Qué NO hace: no inventa proveedores, stock mínimo, equipos ni porcentajes de reparto; no toca el
Excel de comisiones (contiene datos sensibles) ni lo abre bajo ninguna circunstancia.

Cada decisión no trivial (medida en litros, carácter corregido, sin medida...) va al reporte:
consola + CSV en `data/reportes/`.
"""

import argparse
import csv
import sys
import uuid
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.database import get_sessionmaker
from app.db.seeds.catalogo_parser import DescripcionProducto, interpretar_descripcion
from app.domain.models.catalog import (
    Categoria,
    Inventario,
    Producto,
    ProductoAliasExcel,
    Ruta,
)
from app.services.etl_lectura import normalizar_texto, sku_desde_medida_sabor

HOJA_PRODUCTOS = "CODIGOS Y PRECIOS"
HOJAS_RUTA = ("Cornelio", "Antonio", "Cesar")
HOJAS_IGNORADAS = frozenset(
    {"consumo", "luis", "codigos_y_precios", "clientes", "faltante", "merma"}
)
UNIDAD_PAQUETE = "paquete"
MAX_ALIAS = 80
CENTAVOS = Decimal("0.01")


# ------------------------------------------------------------------ modelo de reporte
@dataclass
class FilaProducto:
    fila: int
    sku: str
    precio: Decimal
    descripcion: DescripcionProducto


@dataclass
class Reporte:
    aplicado: bool = False
    productos_creados: list[str] = field(default_factory=list)
    productos_actualizados: list[str] = field(default_factory=list)
    productos_sin_cambios: list[str] = field(default_factory=list)
    categorias_creadas: list[str] = field(default_factory=list)
    rutas_creadas: list[str] = field(default_factory=list)
    rutas_existentes: list[str] = field(default_factory=list)
    alias_creados: list[str] = field(default_factory=list)
    alias_sin_cambios: list[str] = field(default_factory=list)
    # (referencia, detalle)
    no_interpretables: list[tuple[str, str]] = field(default_factory=list)
    duplicados: list[tuple[str, str]] = field(default_factory=list)
    avisos_producto: list[tuple[str, str]] = field(default_factory=list)
    sin_empate: list[tuple[str, str]] = field(default_factory=list)
    alias_ambiguos: list[tuple[str, str]] = field(default_factory=list)
    alias_invalidos: list[tuple[str, str]] = field(default_factory=list)
    hojas_ignoradas: list[str] = field(default_factory=list)
    hojas_no_clasificadas: list[str] = field(default_factory=list)
    notas: list[str] = field(default_factory=list)

    def secciones(self) -> list[tuple[str, list[tuple[str, str]]]]:
        """`(sección, [(referencia, detalle)])` en el orden del CSV."""

        def simples(lista: list[str]) -> list[tuple[str, str]]:
            return [(x, "") for x in lista]

        return [
            ("productos_creados", simples(self.productos_creados)),
            ("productos_actualizados", simples(self.productos_actualizados)),
            ("productos_sin_cambios", simples(self.productos_sin_cambios)),
            ("categorias_creadas", simples(self.categorias_creadas)),
            ("rutas_creadas", simples(self.rutas_creadas)),
            ("rutas_existentes", simples(self.rutas_existentes)),
            ("alias_creados", simples(self.alias_creados)),
            ("alias_sin_cambios", simples(self.alias_sin_cambios)),
            ("filas_no_interpretables", self.no_interpretables),
            ("duplicados", self.duplicados),
            ("avisos_de_interpretacion", self.avisos_producto),
            ("ventas_sin_empate", self.sin_empate),
            ("alias_ambiguos", self.alias_ambiguos),
            ("alias_invalidos", self.alias_invalidos),
            ("hojas_ignoradas", simples(self.hojas_ignoradas)),
            ("hojas_no_clasificadas", simples(self.hojas_no_clasificadas)),
            ("notas", simples(self.notas)),
        ]


# ------------------------------------------------------------------ lectura de Excel
def _abrir(ruta: Path) -> openpyxl.Workbook:
    if "comision" in ruta.name.lower():
        raise ValueError("El Excel de comisiones contiene datos sensibles: no se importa.")
    return openpyxl.load_workbook(ruta, read_only=True, data_only=True)


def _codigo(valor: object) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return str(int(valor)) if float(valor).is_integer() else None
    texto = str(valor).strip().upper()
    return texto or None


def _precio(valor: object) -> Decimal | None:
    if valor is None or isinstance(valor, bool):
        return None
    try:
        precio = Decimal(str(valor)).quantize(CENTAVOS)
    except InvalidOperation:
        return None
    return precio if precio >= 0 else None


def leer_productos(libro: openpyxl.Workbook, reporte: Reporte) -> list[FilaProducto]:
    if HOJA_PRODUCTOS not in libro.sheetnames:
        raise ValueError(f"El libro no tiene la hoja «{HOJA_PRODUCTOS}».")
    filas = libro[HOJA_PRODUCTOS].iter_rows(values_only=True)
    cabecera = next(filas, ())
    columnas = {normalizar_texto(c): i for i, c in enumerate(cabecera) if c is not None}
    faltan = [c for c in ("codigo", "descripcion", "precio") if c not in columnas]
    if faltan:
        raise ValueError(f"Faltan columnas en «{HOJA_PRODUCTOS}»: {', '.join(faltan)}.")

    productos: dict[str, FilaProducto] = {}
    for numero, fila in enumerate(filas, start=2):
        if all(c is None or str(c).strip() == "" for c in fila):
            continue
        referencia = f"{HOJA_PRODUCTOS}!{numero}"
        sku = _codigo(fila[columnas["codigo"]])
        texto = fila[columnas["descripcion"]]
        precio = _precio(fila[columnas["precio"]])
        if sku is None:
            reporte.no_interpretables.append((referencia, "código ausente o no numérico"))
            continue
        if not isinstance(texto, str) or not texto.strip():
            reporte.no_interpretables.append((referencia, f"{sku}: descripción ausente"))
            continue
        if precio is None:
            reporte.no_interpretables.append((referencia, f"{sku}: precio ausente o inválido"))
            continue
        if sku in productos:
            reporte.duplicados.append(
                (referencia, f"código {sku} repetido (se conserva el primero)")
            )
            continue
        descripcion = interpretar_descripcion(texto)
        for aviso in descripcion.avisos:
            reporte.avisos_producto.append((sku, aviso))
        productos[sku] = FilaProducto(numero, sku, precio, descripcion)
    return list(productos.values())


def reportar_duplicados_de_descripcion(filas: list[FilaProducto], reporte: Reporte) -> None:
    """Distintos códigos con la misma presentación (variantes «MC»): se listan, no se unen."""
    grupos: dict[tuple, list[str]] = {}
    for f in filas:
        d = f.descripcion
        clave = (d.marca, d.sabor, d.medida_ml, d.unidades_por_paquete)
        grupos.setdefault(clave, []).append(f.sku)
    for (marca, sabor, medida, unidades), skus in grupos.items():
        if len(skus) > 1:
            reporte.duplicados.append(
                (
                    ", ".join(skus),
                    f"misma presentación: {marca} {sabor} {medida} ml x{unidades}",
                )
            )


def clasificar_hojas(libro: openpyxl.Workbook, reporte: Reporte) -> list[str]:
    """Hojas de vendedor (rutas). Solo las de `HOJAS_RUTA`; el resto se ignora o se reporta."""
    permitidas = {normalizar_texto(h): h for h in HOJAS_RUTA}
    rutas: list[str] = []
    for nombre in libro.sheetnames:
        clave = normalizar_texto(nombre)
        if clave in permitidas:
            rutas.append(nombre)
        elif clave in HOJAS_IGNORADAS:
            reporte.hojas_ignoradas.append(nombre)
        else:
            reporte.hojas_no_clasificadas.append(nombre)
    return rutas


def leer_claves_ventas(libro: openpyxl.Workbook, hojas: list[str]) -> dict[str, set[str]]:
    """`clave <medida>-<SABOR> -> hojas donde aparece` (mismo criterio que el ETL de formato ancho).

    Solo se leen las columnas `Articulo` y `Descripcion` (las dos primeras).
    """
    claves: dict[str, set[str]] = {}
    for hoja in hojas:
        for medida, sabor in libro[hoja].iter_rows(min_row=2, max_col=2, values_only=True):
            if not isinstance(sabor, str) or not sabor.strip() or isinstance(medida, str):
                continue
            claves.setdefault(sku_desde_medida_sabor(medida, sabor), set()).add(hoja)
    return claves


def leer_alias_csv(ruta: Path, reporte: Reporte) -> dict[str, set[str]]:
    """`alias -> {sku}` desde el CSV de equivalencias (`alias`, `sku`). Omitido si falta."""
    if not ruta.exists():
        reporte.notas.append(f"Sin CSV de alias ({ruta.name}); se omite.")
        return {}
    pares: dict[str, set[str]] = {}
    with ruta.open(newline="", encoding="utf-8-sig") as fh:
        lector = csv.DictReader(fh)
        if not lector.fieldnames or {"alias", "sku"} - {
            c.strip().lower() for c in lector.fieldnames
        }:
            reporte.alias_invalidos.append(
                (ruta.name, "el CSV debe tener las columnas alias y sku")
            )
            return {}
        for numero, fila in enumerate(lector, start=2):
            datos = {k.strip().lower(): (v or "").strip() for k, v in fila.items() if k}
            alias, sku = datos.get("alias", "").upper(), datos.get("sku", "").upper()
            referencia = f"{ruta.name}!{numero}"
            if not alias or not sku:
                reporte.alias_invalidos.append((referencia, "alias o sku vacío"))
            elif len(alias) > MAX_ALIAS:
                reporte.alias_invalidos.append(
                    (referencia, f"alias de más de {MAX_ALIAS} caracteres")
                )
            else:
                pares.setdefault(alias, set()).add(sku)
    return pares


# ------------------------------------------------------------------ persistencia
def _categorias(db: Session) -> dict[str, Categoria]:
    return {c.nombre.lower(): c for c in db.scalars(select(Categoria))}


def _cambios(producto: Producto, fila: FilaProducto) -> dict[str, object]:
    """Campos del importador que difieren de lo guardado (no toca categoría, costo ni stock)."""
    d = fila.descripcion
    deseado = {
        "nombre": d.nombre,
        "precio_venta": fila.precio,
        "unidad_medida": UNIDAD_PAQUETE,
        "unidades_por_paquete": d.unidades_por_paquete,
        "medida_ml": d.medida_ml,
        "sabor": d.sabor,
    }
    return {k: v for k, v in deseado.items() if getattr(producto, k) != v}


def _sincronizar_productos(
    db: Session, filas: list[FilaProducto], aplicar: bool, reporte: Reporte
) -> dict[str, uuid.UUID | None]:
    """Upsert por SKU. Devuelve `SKU -> id` (None si se crearía en un dry-run)."""
    existentes = {p.sku.upper(): p for p in db.scalars(select(Producto))}
    categorias = _categorias(db)
    planificadas: set[str] = set()
    ids: dict[str, uuid.UUID | None] = {sku: p.id for sku, p in existentes.items()}

    for fila in filas:
        producto = existentes.get(fila.sku)
        if producto is not None:
            cambios = _cambios(producto, fila)
            if cambios:
                reporte.productos_actualizados.append(fila.sku)
                if aplicar:
                    for campo, valor in cambios.items():
                        setattr(producto, campo, valor)
            else:
                reporte.productos_sin_cambios.append(fila.sku)
            continue

        marca = (fila.descripcion.marca or "Sin marca").title()
        categoria = categorias.get(marca.lower())
        if categoria is None and marca.lower() not in planificadas:
            reporte.categorias_creadas.append(marca)
            planificadas.add(marca.lower())
            if aplicar:
                categoria = Categoria(nombre=marca)
                db.add(categoria)
                db.flush()
                categorias[marca.lower()] = categoria
        reporte.productos_creados.append(fila.sku)
        if aplicar:
            assert categoria is not None
            d = fila.descripcion
            nuevo = Producto(
                sku=fila.sku,
                nombre=d.nombre,
                categoria_id=categoria.id,
                unidad_medida=UNIDAD_PAQUETE,
                precio_venta=fila.precio,
                unidades_por_paquete=d.unidades_por_paquete,
                medida_ml=d.medida_ml,
                sabor=d.sabor,
            )
            db.add(nuevo)
            db.flush()
            db.add(Inventario(producto_id=nuevo.id))  # existencia en 0, como el alta manual
            db.flush()
            ids[fila.sku] = nuevo.id
        else:
            ids[fila.sku] = None
    return ids


def _sincronizar_rutas(db: Session, hojas: list[str], aplicar: bool, reporte: Reporte) -> None:
    """Una ruta por hoja de vendedor: `codigo` = nombre en mayúsculas. No asigna vendedor."""
    for hoja in hojas:
        codigo = hoja.strip().upper()
        existente = db.scalar(select(Ruta).where(func.upper(Ruta.codigo) == codigo))
        if existente is not None:
            reporte.rutas_existentes.append(codigo)
            continue
        reporte.rutas_creadas.append(codigo)
        if aplicar:
            db.add(Ruta(codigo=codigo, nombre=hoja.strip()))
            db.flush()


def _sincronizar_alias(
    db: Session,
    pares: dict[str, set[str]],
    ids: dict[str, uuid.UUID | None],
    aplicar: bool,
    reporte: Reporte,
) -> dict[str, str]:
    """Inserta alias válidos. Devuelve `alias -> sku` de los alias vigentes (BD + CSV válido)."""
    en_bd = {a.alias: a.producto_id for a in db.scalars(select(ProductoAliasExcel))}
    skus_por_id = {pid: sku for sku, pid in ids.items() if pid is not None}
    vigentes = {alias: skus_por_id.get(pid, "?") for alias, pid in en_bd.items()}

    for alias, skus in sorted(pares.items()):
        if len(skus) > 1:
            reporte.alias_ambiguos.append(
                (alias, f"apunta a más de un producto: {', '.join(sorted(skus))}")
            )
            continue
        sku = next(iter(skus))
        if sku not in ids:
            reporte.alias_invalidos.append((alias, f"el SKU {sku} no existe en el catálogo"))
            continue
        producto_id = ids[sku]
        if alias in en_bd:
            if producto_id is not None and en_bd[alias] == producto_id:
                reporte.alias_sin_cambios.append(alias)
            else:
                reporte.alias_ambiguos.append(
                    (alias, f"ya existe apuntando a otro producto; el CSV pide {sku}")
                )
            continue
        reporte.alias_creados.append(alias)
        vigentes[alias] = sku
        if aplicar:
            assert producto_id is not None
            db.add(ProductoAliasExcel(producto_id=producto_id, alias=alias))
            db.flush()
    return vigentes


# ------------------------------------------------------------------ orquestación
def ejecutar(
    db: Session,
    archivo: Path,
    alias_csv: Path | None,
    *,
    aplicar: bool,
) -> Reporte:
    """Corre el importador. Con `aplicar=False` no escribe nada (solo lecturas)."""
    reporte = Reporte(aplicado=aplicar)
    libro = _abrir(archivo)
    try:
        filas = leer_productos(libro, reporte)
        reportar_duplicados_de_descripcion(filas, reporte)
        hojas = clasificar_hojas(libro, reporte)
        claves = leer_claves_ventas(libro, hojas)
    finally:
        libro.close()

    ids = _sincronizar_productos(db, filas, aplicar, reporte)
    _sincronizar_rutas(db, hojas, aplicar, reporte)
    pares = leer_alias_csv(alias_csv, reporte) if alias_csv is not None else {}
    alias_vigentes = _sincronizar_alias(db, pares, ids, aplicar, reporte)

    conocidos = {s.upper() for s in ids} | {a.upper() for a in alias_vigentes}
    for clave in sorted(claves):
        if clave.upper() not in conocidos:
            reporte.sin_empate.append((clave, "hojas: " + ", ".join(sorted(claves[clave]))))

    if aplicar:
        db.commit()
    else:
        db.rollback()
    return reporte


def escribir_csv(reporte: Reporte, carpeta: Path) -> Path:
    carpeta.mkdir(parents=True, exist_ok=True)
    marca = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    destino = carpeta / f"importar_catalogos_{marca}.csv"
    with destino.open("w", newline="", encoding="utf-8-sig") as fh:
        escritor = csv.writer(fh)
        escritor.writerow(["modo", "seccion", "referencia", "detalle"])
        modo = "aplicado" if reporte.aplicado else "dry-run"
        for seccion, filas in reporte.secciones():
            for referencia, detalle in filas:
                escritor.writerow([modo, seccion, referencia, detalle])
    return destino


def imprimir(reporte: Reporte, destino: Path | None) -> None:
    print("=== Importador de catálogos ===")
    print("Modo:", "APLICADO (escribió en la BD)" if reporte.aplicado else "DRY-RUN (no escribió)")
    for seccion, filas in reporte.secciones():
        if filas:
            print(f"- {seccion}: {len(filas)}")
    for seccion in ("filas_no_interpretables", "duplicados", "alias_ambiguos", "alias_invalidos"):
        for nombre, filas in reporte.secciones():
            if nombre == seccion and filas:
                print(f"\n[{seccion}]")
                for referencia, detalle in filas[:20]:
                    print(f"  {referencia}: {detalle}")
                if len(filas) > 20:
                    print(f"  ... y {len(filas) - 20} más (ver CSV)")
    if destino is not None:
        print(f"\nReporte completo: {destino}")


def _archivo_ventas(raw: Path) -> Path:
    candidatos = sorted(raw.glob("VENTAS DIARIAS*.xlsx"))
    if not candidatos:
        raise FileNotFoundError(f"No hay un libro «VENTAS DIARIAS*.xlsx» en {raw}.")
    return candidatos[-1]


def main(argv: list[str] | None = None) -> int:
    datos = get_settings().data_dir
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--aplicar", action="store_true", help="Escribe en la BD (por defecto, dry-run)."
    )
    parser.add_argument(
        "--archivo", type=Path, help="Libro de ventas diarias (por defecto data/raw)."
    )
    parser.add_argument(
        "--alias-csv",
        type=Path,
        default=datos / "plantillas" / "alias_excel.csv",
        help="CSV de equivalencias (alias, sku); se omite si no existe.",
    )
    parser.add_argument("--reportes", type=Path, default=datos / "reportes")
    args = parser.parse_args(argv)

    archivo = args.archivo or _archivo_ventas(datos / "raw")
    sesion = get_sessionmaker()()
    try:
        reporte = ejecutar(sesion, archivo, args.alias_csv, aplicar=args.aplicar)
    except Exception:
        sesion.rollback()
        raise
    finally:
        sesion.close()
    destino = escribir_csv(reporte, args.reportes)
    imprimir(reporte, destino)
    return 0


if __name__ == "__main__":
    sys.exit(main())
