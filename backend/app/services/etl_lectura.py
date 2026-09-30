"""Lectura y normalización de los Excel de ventas (sin acceso a BD).

Soporta dos formatos y ambos producen `FilaVenta`, de modo que la validación contra el
catálogo y la persistencia (`etl_service`) sean únicas:

* **tabular** (plantilla del contrato, TC-ETL-*): una fila por venta con las columnas
  `fecha, ruta, sku, cantidad, precio_unitario` y `vendedor` opcional.
* **ancho** (Excel comercial real, `VENTAS DIARIAS`): una hoja por vendedor/ruta; producto
  por *medida* + *sabor* y un bloque de columnas por día (`ventas NN`).

Números de fila: son los de Excel (la cabecera es la fila 1).
"""

import calendar
import re
import unicodedata
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.schemas.etl import CodigoErrorEtl

FORMATO_TABULAR = "tabular"
FORMATO_ANCHO = "ancho"

# Alias de cabecera (normalizada) -> columna canónica del formato tabular.
ALIAS_COLUMNAS: dict[str, tuple[str, ...]] = {
    "fecha": ("fecha", "fecha_venta"),
    "ruta": ("ruta", "codigo_ruta", "ruta_codigo", "cod_ruta"),
    "sku": ("sku", "codigo", "codigo_sku", "codigo_producto"),
    "cantidad": ("cantidad", "unidades", "cant"),
    "precio_unitario": ("precio_unitario", "precio", "precio_venta"),
    "vendedor": ("vendedor", "vendedor_email", "email_vendedor"),
}
COLUMNAS_OBLIGATORIAS = ("fecha", "ruta", "sku", "cantidad", "precio_unitario")

_MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}  # fmt: skip
_FORMATOS_FECHA = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y")
_LIMITE_CANTIDAD = Decimal(10) ** 10  # numeric(12,2)
_LIMITE_MONTO = Decimal(10) ** 12  # numeric(14,2)
_CENTAVOS = Decimal("0.01")
_RE_DIA_VENTAS = re.compile(r"^ventas_(\d{1,2})$")


class ArchivoInvalido(Exception):
    """El archivo no se puede procesar en absoluto (no es .xlsx, hoja inexistente, etc.)."""

    def __init__(self, codigo: str, mensaje: str) -> None:
        self.codigo = codigo
        self.mensaje = mensaje
        super().__init__(mensaje)


@dataclass(frozen=True)
class ErrorLectura:
    fila: int
    columna: str
    mensaje: str
    categoria: CodigoErrorEtl
    valor: str | None = None
    hoja: str | None = None
    # Distingue registros de una misma fila de Excel (formato ancho: uno por día).
    registro: str = ""


@dataclass(frozen=True)
class FilaVenta:
    hoja: str
    fila: int
    fecha: date
    ruta: str
    sku: str
    cantidad: Decimal
    precio_unitario: Decimal
    vendedor: str | None = None
    registro: str = ""

    @property
    def monto_total(self) -> Decimal:
        return (self.cantidad * self.precio_unitario).quantize(_CENTAVOS, ROUND_HALF_UP)


@dataclass
class LecturaExcel:
    formato: str
    filas: list[FilaVenta] = field(default_factory=list)
    errores: list[ErrorLectura] = field(default_factory=list)
    # Filas de datos leídas (con o sin error); base de `filas_totales` del lote.
    filas_totales: int = 0
    # Identificador (hoja, fila, registro) de los registros con al menos un error de lectura.
    filas_con_error: set[tuple[str, int, str]] = field(default_factory=set)
    columnas_faltantes: list[str] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)


# ------------------------------------------------------------------ utilidades de celda
def normalizar_texto(valor: object) -> str:
    """Minúsculas, sin acentos y con `_` en lugar de espacios/símbolos (para cabeceras)."""
    texto = unicodedata.normalize("NFKD", str(valor if valor is not None else ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", "_", texto).strip("_")


def _texto_visible(valor: Any) -> str | None:
    return None if valor is None else str(valor)


def parsear_fecha(valor: Any) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        for formato in _FORMATOS_FECHA:
            try:
                return datetime.strptime(valor.strip(), formato).date()
            except ValueError:
                continue
    return None


def parsear_decimal(valor: Any) -> Decimal | None:
    """Convierte a `Decimal` con 2 decimales; `None` si no es un número finito."""
    if valor is None or isinstance(valor, bool):
        return None
    try:
        numero = Decimal(str(valor).strip())
    except InvalidOperation:
        return None
    if not numero.is_finite():
        return None
    return numero.quantize(_CENTAVOS, ROUND_HALF_UP)


def _fila_vacia(fila: tuple[Any, ...]) -> bool:
    return all(c is None or (isinstance(c, str) and not c.strip()) for c in fila)


def _celda(fila: tuple[Any, ...], indice: int | None) -> Any:
    return fila[indice] if indice is not None and indice < len(fila) else None


def periodo_desde_nombre(nombre_archivo: str) -> str | None:
    """`VENTAS DIARIAS MARZO 2025.xlsx` -> `2025-03`."""
    texto = normalizar_texto(nombre_archivo)
    anio = re.search(r"(?<!\d)(20\d{2})(?!\d)", texto)
    # `normalizar_texto` deja `_` como separador y `\b` no lo trata como límite de palabra.
    mes = next(
        (n for nombre, n in _MESES.items() if re.search(rf"(?<![a-z]){nombre}(?![a-z])", texto)),
        None,
    )
    if anio and mes:
        return f"{anio.group(1)}-{mes:02d}"
    return None


# ------------------------------------------------------------------ punto de entrada
def leer_excel(
    contenido: bytes,
    *,
    nombre_archivo: str,
    hoja: str | None = None,
    periodo: str | None = None,
) -> LecturaExcel:
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except (InvalidFileException, OSError, KeyError, ValueError) as exc:
        raise ArchivoInvalido("ARCHIVO_INVALIDO", "El archivo no es un .xlsx válido.") from exc
    except Exception as exc:  # zipfile.BadZipFile u otros errores de parseo del contenedor
        raise ArchivoInvalido("ARCHIVO_INVALIDO", "El archivo no es un .xlsx válido.") from exc

    try:
        if hoja is not None and hoja not in libro.sheetnames:
            raise ArchivoInvalido("HOJA_NO_ENCONTRADA", f"La hoja '{hoja}' no existe en el libro.")
        candidatas = [hoja] if hoja is not None else list(libro.sheetnames)
        if not candidatas:
            raise ArchivoInvalido("ARCHIVO_INVALIDO", "El libro no contiene hojas.")

        anchas: list[tuple[str, dict[str, Any]]] = []
        omitidas: list[str] = []
        primera: tuple[str, tuple[Any, ...], Iterator[tuple[Any, ...]]] | None = None
        for nombre in candidatas:
            filas = libro[nombre].iter_rows(values_only=True)
            cabecera = next(filas, None)
            if primera is None:
                primera = (nombre, cabecera or (), filas)
            layout = _detectar_layout_ancho(cabecera or ())
            if layout is not None:
                anchas.append((nombre, layout))
            else:
                omitidas.append(nombre)

        if anchas:
            return _leer_ancho(libro, anchas, omitidas, nombre_archivo, hoja, periodo)
        assert primera is not None
        return _leer_tabular(*primera)
    finally:
        libro.close()


# ------------------------------------------------------------------ formato tabular
def _mapear_cabecera(cabecera: tuple[Any, ...]) -> dict[str, int]:
    normalizadas = {normalizar_texto(c): i for i, c in enumerate(cabecera) if c is not None}
    columnas: dict[str, int] = {}
    for canonica, alias in ALIAS_COLUMNAS.items():
        encontrado = next((a for a in alias if a in normalizadas), None)
        if encontrado is not None:
            columnas[canonica] = normalizadas[encontrado]
    return columnas


def _leer_tabular(
    hoja: str, cabecera: tuple[Any, ...], filas: Iterator[tuple[Any, ...]]
) -> LecturaExcel:
    lectura = LecturaExcel(formato=FORMATO_TABULAR)
    columnas = _mapear_cabecera(cabecera)
    lectura.columnas_faltantes = [c for c in COLUMNAS_OBLIGATORIAS if c not in columnas]
    for faltante in lectura.columnas_faltantes:
        lectura.errores.append(
            ErrorLectura(
                fila=1,
                columna=faltante,
                mensaje=f"Falta la columna obligatoria '{faltante}'.",
                categoria=CodigoErrorEtl.COLUMNAS_FALTANTES,
                hoja=hoja,
            )
        )

    for numero, fila in enumerate(filas, start=2):
        if _fila_vacia(fila):
            continue
        lectura.filas_totales += 1
        errores = _validar_fila_tabular(hoja, numero, fila, columnas)
        if errores:
            lectura.errores.extend(errores)
            lectura.filas_con_error.add((hoja, numero, ""))
            continue
        if lectura.columnas_faltantes:
            continue  # sin todas las columnas no se puede construir la venta
        vendedor = _celda(fila, columnas.get("vendedor"))
        lectura.filas.append(
            FilaVenta(
                hoja=hoja,
                fila=numero,
                fecha=parsear_fecha(_celda(fila, columnas["fecha"])),  # type: ignore[arg-type]
                ruta=str(_celda(fila, columnas["ruta"])).strip(),
                sku=str(_celda(fila, columnas["sku"])).strip(),
                cantidad=parsear_decimal(_celda(fila, columnas["cantidad"])),  # type: ignore[arg-type]
                precio_unitario=parsear_decimal(_celda(fila, columnas["precio_unitario"])),  # type: ignore[arg-type]
                vendedor=str(vendedor).strip() if vendedor not in (None, "") else None,
            )
        )
    return lectura


def _validar_fila_tabular(
    hoja: str, numero: int, fila: tuple[Any, ...], columnas: dict[str, int]
) -> list[ErrorLectura]:
    errores: list[ErrorLectura] = []

    def error(columna: str, mensaje: str, categoria: CodigoErrorEtl, valor: Any) -> None:
        errores.append(
            ErrorLectura(numero, columna, mensaje, categoria, _texto_visible(valor), hoja)
        )

    if "fecha" in columnas:
        valor = _celda(fila, columnas["fecha"])
        if parsear_fecha(valor) is None:
            error("fecha", "Fecha inválida o ausente.", CodigoErrorEtl.TIPOS_INVALIDOS, valor)

    for columna in ("ruta", "sku"):
        if columna in columnas:
            valor = _celda(fila, columnas[columna])
            if valor is None or not str(valor).strip():
                error(
                    columna, f"'{columna}' es obligatorio.", CodigoErrorEtl.DATOS_INVALIDOS, valor
                )

    numericos = {}
    for columna, limite in (("cantidad", _LIMITE_CANTIDAD), ("precio_unitario", _LIMITE_MONTO)):
        if columna not in columnas:
            continue
        valor = _celda(fila, columnas[columna])
        numero_dec = parsear_decimal(valor)
        if numero_dec is None:
            error(
                columna, "Valor numérico inválido o ausente.", CodigoErrorEtl.TIPOS_INVALIDOS, valor
            )
        elif numero_dec < 0:
            error(columna, "No puede ser negativo.", CodigoErrorEtl.DATOS_INVALIDOS, valor)
        elif numero_dec >= limite:
            error(columna, "Excede el máximo permitido.", CodigoErrorEtl.DATOS_INVALIDOS, valor)
        else:
            numericos[columna] = numero_dec

    if (
        len(numericos) == 2
        and numericos["cantidad"] * numericos["precio_unitario"] >= _LIMITE_MONTO
    ):
        error(
            "monto_total",
            "El monto (cantidad × precio) excede el máximo permitido.",
            CodigoErrorEtl.DATOS_INVALIDOS,
            None,
        )
    return errores


# ------------------------------------------------------------------ formato ancho
def _detectar_layout_ancho(cabecera: tuple[Any, ...]) -> dict[str, Any] | None:
    """Índices de las columnas del layout `VENTAS DIARIAS`, o `None` si la hoja no lo cumple."""
    indices: dict[str, int] = {}
    dias: dict[int, int] = {}
    for i, c in enumerate(cabecera):
        clave = normalizar_texto(c)
        if clave in ("articulo", "descripcion", "precio") and clave not in indices:
            indices[clave] = i
        elif (m := _RE_DIA_VENTAS.match(clave)) and int(m.group(1)) not in dias:
            dias[int(m.group(1))] = i
    if {"articulo", "descripcion", "precio"} <= indices.keys() and dias:
        return {**indices, "dias": dias}
    return None


def _leer_ancho(
    libro: Any,
    hojas: list[tuple[str, dict[str, Any]]],
    omitidas: list[str],
    nombre_archivo: str,
    hoja_solicitada: str | None,
    periodo: str | None,
) -> LecturaExcel:
    periodo = periodo or periodo_desde_nombre(nombre_archivo)
    if periodo is None:
        raise ArchivoInvalido(
            "PERIODO_NO_DETECTADO",
            "No se pudo inferir el mes del archivo; envíe el campo 'periodo' (YYYY-MM).",
        )
    anio, mes = (int(p) for p in periodo.split("-"))
    dias_del_mes = calendar.monthrange(anio, mes)[1]

    lectura = LecturaExcel(formato=FORMATO_ANCHO)
    if hoja_solicitada is None and omitidas:
        lectura.advertencias.append(
            "Hojas omitidas por no tener el layout de ventas diarias: " + ", ".join(omitidas) + "."
        )

    for nombre, layout in hojas:
        filas = libro[nombre].iter_rows(values_only=True)
        next(filas, None)  # cabecera
        for numero, fila in enumerate(filas, start=2):
            medida, sabor = _celda(fila, layout["articulo"]), _celda(fila, layout["descripcion"])
            # Filtra totales, metadatos ("Nit", "Digifact"...) y filas sin producto.
            if not isinstance(sabor, str) or not sabor.strip() or isinstance(medida, str):
                continue
            precio = parsear_decimal(_celda(fila, layout["precio"]))
            sku = sku_desde_medida_sabor(medida, sabor)
            for dia, indice in sorted(layout["dias"].items()):
                valor = _celda(fila, indice)
                cantidad = parsear_decimal(valor)
                if valor is None or cantidad == 0 or (isinstance(valor, str) and not valor.strip()):
                    continue
                lectura.filas_totales += 1
                columna = f"ventas {dia:02d}"
                categoria, mensaje = None, None
                if cantidad is None:
                    categoria, mensaje = CodigoErrorEtl.TIPOS_INVALIDOS, "Cantidad no numérica."
                elif cantidad < 0 or cantidad >= _LIMITE_CANTIDAD:
                    categoria, mensaje = CodigoErrorEtl.DATOS_INVALIDOS, "Cantidad fuera de rango."
                elif dia > dias_del_mes:
                    categoria = CodigoErrorEtl.DATOS_INVALIDOS
                    mensaje = f"El día {dia:02d} no existe en {periodo}."
                elif precio is None or precio < 0 or precio >= _LIMITE_MONTO:
                    categoria, mensaje = (
                        CodigoErrorEtl.DATOS_INVALIDOS,
                        "Precio ausente o inválido.",
                    )
                if categoria is not None:
                    lectura.errores.append(
                        ErrorLectura(
                            numero,
                            columna,
                            mensaje,
                            categoria,
                            _texto_visible(valor),
                            nombre,
                            registro=columna,
                        )
                    )
                    lectura.filas_con_error.add((nombre, numero, columna))
                    continue
                lectura.filas.append(
                    FilaVenta(
                        hoja=nombre,
                        fila=numero,
                        fecha=date(anio, mes, dia),
                        ruta=nombre,
                        sku=sku,
                        cantidad=cantidad,  # type: ignore[arg-type]
                        precio_unitario=precio,  # type: ignore[arg-type]
                        registro=columna,
                    )
                )
    return lectura


def sku_desde_medida_sabor(medida: Any, sabor: str) -> str:
    """Clave de producto del Excel comercial (no trae SKU): `<medida>-<SABOR>` normalizado.

    Ej.: (3030, 'Piña') -> '3030-PINA'; (None, 'PULPIN PERA') -> 'SM-PULPIN_PERA'.
    El catálogo (`productos.sku`) debe cargarse con esta misma convención para el formato ancho.
    """
    if medida is None:
        texto_medida = "SM"
    else:
        numero = parsear_decimal(medida)
        texto_medida = str(medida) if numero is None else f"{numero.normalize():f}"
    return f"{texto_medida}-{normalizar_texto(sabor).upper()}"
