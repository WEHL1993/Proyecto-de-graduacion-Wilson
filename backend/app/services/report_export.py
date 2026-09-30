"""Serialización de reportes a CSV y Excel. Sin acceso a datos: recibe tablas ya calculadas."""

import csv
import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

Celda = str | int | float | Decimal | date | None

MEDIA_TYPE_CSV = "text/csv; charset=utf-8"
MEDIA_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class Tabla:
    nombre: str  # título de la hoja (máx. 31 caracteres en Excel)
    columnas: list[str]
    filas: list[list[Celda]]


def _texto(valor: Celda) -> str:
    return "" if valor is None else str(valor)


def a_csv(tablas: list[Tabla]) -> bytes:
    """UTF-8 con BOM (Excel lo abre con tildes). Varias tablas se separan por una línea en
    blanco y un título de sección."""
    salida = io.StringIO()
    escritor = csv.writer(salida, lineterminator="\r\n")
    for i, tabla in enumerate(tablas):
        if i:
            escritor.writerow([])
        if len(tablas) > 1:
            escritor.writerow([tabla.nombre])
        escritor.writerow(tabla.columnas)
        for fila in tabla.filas:
            escritor.writerow([_texto(c) for c in fila])
    return salida.getvalue().encode("utf-8-sig")


def a_xlsx(tablas: list[Tabla]) -> bytes:
    libro = Workbook()
    libro.remove(libro.active)
    for tabla in tablas:
        hoja = libro.create_sheet(tabla.nombre[:31])
        hoja.append(tabla.columnas)
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        for fila in tabla.filas:
            # Decimal se escribe como número (no texto) para que Excel pueda operarlo.
            hoja.append([float(c) if isinstance(c, Decimal) else c for c in fila])
        for indice, columna in enumerate(tabla.columnas, start=1):
            ancho = max([len(columna), *(len(_texto(f[indice - 1])) for f in tabla.filas)])
            hoja.column_dimensions[get_column_letter(indice)].width = min(ancho + 2, 48)
        hoja.freeze_panes = "A2"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
